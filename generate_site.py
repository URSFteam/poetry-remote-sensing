from __future__ import annotations

import hashlib
import html
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.run import Run


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT.parent / "诗词遥感-草稿 - 小修版-诗词楷体版.docx"
SITE = ROOT / "site"
ASSETS = SITE / "assets"
MEDIA = ASSETS / "media"

PAGES = [
    (5, "preface.html", "作者自序"),
    (19, "chapter-1.html", "一、遥感与诗词的交融"),
    (52, "chapter-2.html", "二、诗词中的辐射光源"),
    (102, "chapter-3.html", "三、诗词中的大气作用"),
    (155, "chapter-4.html", "四、诗词中的地表结构"),
    (238, "chapter-5.html", "五、诗词中的地物反射"),
    (275, "chapter-6.html", "六、诗词中的遥感观测"),
    (329, "chapter-7.html", "七、诗词中的遥感建模"),
    (369, "chapter-8.html", "八、诗词中的遥感解译"),
    (406, "appendix.html", "作者的两三首词摘录"),
    (453, "references.html", "参考资料"),
]


def reset_output() -> None:
    if SITE.exists():
        shutil.rmtree(SITE)
    MEDIA.mkdir(parents=True)


def extension_for(content_type: str, blob: bytes) -> str:
    known = {
        "image/jpeg": ".jpg",
        "image/png": ".png",
        "image/gif": ".gif",
        "image/webp": ".webp",
        "image/svg+xml": ".svg",
        "image/tiff": ".tif",
        "image/bmp": ".bmp",
    }
    if content_type in known:
        return known[content_type]
    if blob.startswith(b"\x89PNG"):
        return ".png"
    if blob.startswith(b"\xff\xd8"):
        return ".jpg"
    if blob.startswith((b"GIF87a", b"GIF89a")):
        return ".gif"
    return ".bin"


def extract_images(doc: Document) -> dict[str, str]:
    result: dict[str, str] = {}
    by_digest: dict[str, str] = {}
    sequence = 1
    for rel_id, rel in doc.part.rels.items():
        if "image" not in rel.reltype:
            continue
        part = rel.target_part
        blob = part.blob
        digest = hashlib.sha256(blob).hexdigest()
        if digest not in by_digest:
            ext = extension_for(part.content_type, blob)
            filename = f"image-{sequence:03d}{ext}"
            (MEDIA / filename).write_bytes(blob)
            by_digest[digest] = filename
            sequence += 1
        result[rel_id] = by_digest[digest]
    return result


def is_kaiti(run: Run) -> bool:
    rpr = run._r.rPr
    if rpr is None:
        return False
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        return False
    names = [
        fonts.get(qn("w:eastAsia"), ""),
        fonts.get(qn("w:ascii"), ""),
        fonts.get(qn("w:hAnsi"), ""),
    ]
    return any("楷" in name or "Kai" in name for name in names)


URL_RE = re.compile(r"(https?://[^\s<>]+)")


def link_urls(text: str) -> str:
    escaped = html.escape(text).replace("\n", "<br>\n")
    return URL_RE.sub(r'<a href="\1">\1</a>', escaped)


def render_run(run: Run) -> str:
    if not run.text:
        return ""
    value = link_urls(run.text)
    if run.bold:
        value = f"<strong>{value}</strong>"
    if run.italic:
        value = f"<em>{value}</em>"
    if run.font.superscript:
        value = f"<sup>{value}</sup>"
    elif run.font.subscript:
        value = f"<sub>{value}</sub>"
    if is_kaiti(run):
        value = f'<span class="poetry">{value}</span>'
    return value


def paragraph_runs(paragraph) -> list[Run]:
    # python-docx's paragraph.runs can omit runs nested in hyperlinks.
    return [Run(node, paragraph) for node in paragraph._p.xpath(".//w:r")]


def paragraph_images(paragraph, image_map: dict[str, str]) -> list[str]:
    filenames: list[str] = []
    for blip in paragraph._p.xpath(".//a:blip"):
        rel_id = blip.get(qn("r:embed"))
        if rel_id and rel_id in image_map:
            filenames.append(image_map[rel_id])
    return filenames


def alignment_class(paragraph) -> str:
    alignment = paragraph.alignment
    if alignment is None:
        return ""
    name = str(alignment).upper()
    if "CENTER" in name:
        return " center"
    if "RIGHT" in name:
        return " right"
    if "JUSTIFY" in name:
        return " justify"
    return ""


def render_paragraph(paragraph, index: int, page_start: int, image_map: dict[str, str]) -> str:
    text_value = paragraph.text.strip()
    runs_html = "".join(render_run(run) for run in paragraph_runs(paragraph)).strip()
    images = paragraph_images(paragraph, image_map)
    blocks: list[str] = []

    if text_value:
        align = alignment_class(paragraph)
        if index == page_start:
            blocks.append(f'<h1 class="page-title{align}">{runs_html}</h1>')
        elif re.match(r"^\d+\.\d+\s+", text_value):
            anchor = f"section-{text_value.split()[0].replace('.', '-')}"
            blocks.append(f'<h2 id="{anchor}" class="section-title{align}">{runs_html}</h2>')
        elif re.match(r"^图\s*\d", text_value):
            blocks.append(f'<p class="caption{align}">{runs_html}</p>')
        else:
            blocks.append(f'<p class="body-text{align}">{runs_html}</p>')

    for filename in images:
        blocks.append(
            '<figure class="document-image">'
            f'<img src="assets/media/{html.escape(filename)}" alt="文中插图" loading="lazy">'
            "</figure>"
        )
    return "\n".join(blocks)


def nav_html(active: str) -> str:
    items = [('index.html', '封面与目录')] + [(filename, title) for _, filename, title in PAGES]
    rows = []
    for filename, title in items:
        current = ' aria-current="page" class="active"' if filename == active else ""
        rows.append(f'<li><a href="{filename}"{current}>{html.escape(title)}</a></li>')
    return "\n".join(rows)


def shell(title: str, active: str, body: str, previous_page: str | None, next_page: str | None) -> str:
    pager = []
    if previous_page:
        pager.append(f'<a class="previous" href="{previous_page}">← 上一页</a>')
    else:
        pager.append("<span></span>")
    if next_page:
        pager.append(f'<a class="next" href="{next_page}">下一页 →</a>')
    return f'''<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta name="description" content="《诗词遥感》电子版">
  <title>{html.escape(title)}｜诗词遥感</title>
  <link rel="stylesheet" href="assets/style.css">
</head>
<body>
  <button class="menu-button" type="button" aria-label="打开目录" aria-expanded="false">目录</button>
  <aside class="sidebar" aria-label="全书目录">
    <a class="book-name" href="index.html">诗词遥感</a>
    <nav><ol>{nav_html(active)}</ol></nav>
  </aside>
  <main class="content">
    <article>{body}</article>
    <nav class="pager" aria-label="翻页">{''.join(pager)}</nav>
    <footer>《诗词遥感》电子版</footer>
  </main>
  <script src="assets/site.js"></script>
</body>
</html>
'''


STYLE = r'''
:root {
  --ink: #25221f;
  --muted: #716a62;
  --paper: #fffef9;
  --sidebar: #f4f0e7;
  --accent: #315c4d;
  --line: #ddd5c7;
}
* { box-sizing: border-box; }
html { scroll-behavior: smooth; }
body {
  margin: 0;
  color: var(--ink);
  background: var(--paper);
  font-family: "Noto Serif SC", "Source Han Serif SC", "Songti SC", SimSun, serif;
  font-size: 18px;
  line-height: 1.95;
}
a { color: var(--accent); text-underline-offset: .18em; }
.sidebar {
  position: fixed;
  inset: 0 auto 0 0;
  width: 310px;
  overflow-y: auto;
  padding: 36px 26px;
  background: var(--sidebar);
  border-right: 1px solid var(--line);
}
.book-name {
  display: block;
  margin: 0 0 24px;
  color: var(--ink);
  font-size: 28px;
  font-weight: 700;
  letter-spacing: .16em;
  text-decoration: none;
}
.sidebar ol { margin: 0; padding: 0; list-style: none; }
.sidebar li { margin: 3px 0; }
.sidebar a:not(.book-name) {
  display: block;
  padding: 7px 10px;
  border-radius: 6px;
  color: #514b45;
  font-size: 14px;
  line-height: 1.5;
  text-decoration: none;
}
.sidebar a:hover, .sidebar a.active { color: white; background: var(--accent); }
.content { max-width: 1050px; margin-left: 310px; padding: 58px 8vw 36px; }
article { max-width: 820px; margin: 0 auto; }
.page-title {
  margin: 0 0 1.2em;
  font-size: clamp(30px, 4vw, 42px);
  line-height: 1.35;
  letter-spacing: .04em;
}
.section-title {
  margin: 2.5em 0 1em;
  padding-bottom: .35em;
  border-bottom: 1px solid var(--line);
  font-size: clamp(23px, 3vw, 30px);
  line-height: 1.45;
}
.body-text { margin: .9em 0; text-indent: 2em; }
.poetry { font-family: "Kaiti SC", STKaiti, KaiTi, "楷体", serif; }
.center { text-align: center; text-indent: 0; }
.right { text-align: right; text-indent: 0; }
.justify { text-align: justify; }
.document-image { margin: 30px auto; text-align: center; }
.document-image img { display: inline-block; max-width: 100%; height: auto; border-radius: 2px; }
.caption { margin: -20px 0 28px; color: var(--muted); font-size: 14px; text-align: center; text-indent: 0; }
.cover { display: grid; min-height: 80vh; place-items: center; text-align: center; }
.cover h1 { margin: 0; font-size: clamp(44px, 8vw, 78px); letter-spacing: .2em; }
.cover p { color: var(--muted); letter-spacing: .12em; text-indent: 0; }
.cover img { max-width: min(100%, 680px); max-height: 70vh; object-fit: contain; }
.toc { margin: 54px auto 20px; }
.toc h2 { font-size: 30px; }
.toc ol { padding-left: 1.4em; }
.toc li { margin: .55em 0; }
.pager { display: flex; justify-content: space-between; max-width: 820px; margin: 64px auto 0; padding-top: 22px; border-top: 1px solid var(--line); }
.pager a { text-decoration: none; }
footer { max-width: 820px; margin: 42px auto 0; color: var(--muted); font-size: 13px; text-align: center; }
.menu-button { display: none; }
@media (max-width: 860px) {
  body { font-size: 17px; }
  .menu-button { display: block; position: fixed; z-index: 20; top: 12px; left: 12px; padding: 8px 12px; border: 1px solid var(--line); border-radius: 6px; background: var(--paper); color: var(--ink); }
  .sidebar { z-index: 10; width: min(84vw, 320px); padding-top: 62px; transform: translateX(-105%); transition: transform .2s ease; box-shadow: 5px 0 22px #0002; }
  body.menu-open .sidebar { transform: translateX(0); }
  .content { margin-left: 0; padding: 72px 22px 32px; }
  .body-text { text-align: justify; }
}
@media print {
  .sidebar, .menu-button, .pager { display: none; }
  .content { max-width: none; margin: 0; padding: 0; }
  body { font-size: 12pt; }
}
'''

SCRIPT = r'''
const button = document.querySelector('.menu-button');
if (button) {
  button.addEventListener('click', () => {
    const open = document.body.classList.toggle('menu-open');
    button.setAttribute('aria-expanded', String(open));
  });
}
'''


def build() -> None:
    reset_output()
    doc = Document(SOURCE)
    image_map = extract_images(doc)

    generated_pages: list[dict[str, object]] = []
    filenames = ["index.html"] + [entry[1] for entry in PAGES]
    for page_no, (start, filename, display_title) in enumerate(PAGES, start=1):
        end = PAGES[page_no][0] if page_no < len(PAGES) else len(doc.paragraphs)
        content = []
        for index in range(start, end):
            block = render_paragraph(doc.paragraphs[index], index, start, image_map)
            if block:
                content.append(block)
        prev_filename = filenames[page_no - 1]
        next_filename = filenames[page_no + 1] if page_no + 1 < len(filenames) else None
        (SITE / filename).write_text(
            shell(display_title, filename, "\n".join(content), prev_filename, next_filename),
            encoding="utf-8",
        )
        generated_pages.append({"file": filename, "start": start, "end": end, "title": display_title})

    cover_images = paragraph_images(doc.paragraphs[0], image_map)
    cover_html = (
        f'<img src="assets/media/{cover_images[0]}" alt="《诗词遥感》封面">'
        if cover_images
        else '<div><h1>诗词遥感</h1><p>电子版</p></div>'
    )
    toc_rows = "\n".join(
        f'<li><a href="{filename}">{html.escape(title)}</a></li>' for _, filename, title in PAGES
    )
    home_body = f'''<section class="cover">{cover_html}</section>
<section class="toc"><h2>目录</h2><ol>{toc_rows}</ol></section>'''
    (SITE / "index.html").write_text(
        shell("封面与目录", "index.html", home_body, None, PAGES[0][1]), encoding="utf-8"
    )
    (ASSETS / "style.css").write_text(STYLE.strip() + "\n", encoding="utf-8")
    (ASSETS / "site.js").write_text(SCRIPT.strip() + "\n", encoding="utf-8")

    report = {
        "source": str(SOURCE),
        "paragraphs": len(doc.paragraphs),
        "image_relationships": len(image_map),
        "unique_images": len(set(image_map.values())),
        "pages": generated_pages,
    }
    (ROOT / "build-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build()
